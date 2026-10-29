from django.db import transaction, IntegrityError
from django.db.models import Q
from django.utils.dateparse import parse_date
from datetime import date, datetime

from openpyxl import load_workbook

from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from common.services import get_user_organization
from app.audit.services import log_activity
from app.attendance.models import Shift
from app.attendance.services import get_employee_for_user

from .models import WeeklyOffRule, CalendarHoliday
from .serializers import (
    WeeklyOffRuleSerializer,
    CalendarHolidaySerializer,
    CalendarHolidayWriteSerializer,
    EmployeeCalendarRequestSerializer,
    EmployeeCalendarDaySerializer,
)
from .services import get_employee_calendar


# =========================================================
# WEEKLY OFF RULE LIST
# =========================================================

class WeeklyOffRuleListAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        rules = WeeklyOffRule.objects.filter(
            shift__organization=organization,
            is_active=True,
        ).select_related("shift")

        shift_id = request.data.get("shift_id")

        if shift_id:
            rules = rules.filter(shift_id=shift_id)

        serializer = WeeklyOffRuleSerializer(
            rules.order_by(
                "shift__name",
                "day_of_week",
                "week_of_month",
                "-effective_from",
            ),
            many=True,
            context={"organization": organization},
        )

        return Response(
            {
                "status": True,
                "message": "Weekly-off rules fetched successfully.",
                "data": serializer.data,
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# WEEKLY OFF RULE CREATE
# =========================================================

class WeeklyOffRuleCreateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        data = request.data.copy()
        data.pop("organization_id", None)

        serializer = WeeklyOffRuleSerializer(
            data=data,
            context={"organization": organization},
        )

        if not serializer.is_valid():
            return Response(
                {
                    "status": False,
                    "message": "Validation failed.",
                    "errors": serializer.errors,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        rule = serializer.save(
            created_by=request.user,
            updated_by=request.user,
        )

        log_activity(
            request=request,
            action="CREATE",
            module="CALENDAR",
            object_type="WeeklyOffRule",
            object_id=rule.id,
            description=(
                f"Weekly-off rule created for shift {rule.shift.name}."
            ),
            new_data={
                "shift_id": rule.shift_id,
                "day_of_week": rule.day_of_week,
                "week_of_month": rule.week_of_month,
                "effective_from": str(rule.effective_from),
                "effective_to": (
                    str(rule.effective_to)
                    if rule.effective_to else None
                ),
            },
        )

        return Response(
            {
                "status": True,
                "message": "Weekly-off rule created successfully.",
                "data": WeeklyOffRuleSerializer(
                    rule,
                    context={"organization": organization},
                ).data,
            },
            status=status.HTTP_201_CREATED,
        )


# =========================================================
# CREATE THE SAME WEEKLY-OFF RULE FOR ALL ACTIVE SHIFTS
# =========================================================

class WeeklyOffRuleBulkCreateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        shifts = list(
            Shift.objects.filter(
                organization=organization,
                is_active=True,
            ).order_by("id")
        )

        if not shifts:
            return Response(
                {
                    "status": False,
                    "message": "No active shifts found.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        data = request.data.copy()
        data.pop("organization_id", None)
        data.pop("shift", None)
        data.pop("shift_id", None)

        # Validate the rule for every shift before creating any rows.
        serializers_to_save = []

        for shift in shifts:
            shift_data = data.copy()
            shift_data["shift"] = shift.id

            serializer = WeeklyOffRuleSerializer(
                data=shift_data,
                context={"organization": organization},
            )

            if not serializer.is_valid():
                return Response(
                    {
                        "status": False,
                        "message": (
                            "Weekly-off rule validation failed. "
                            "No rules were created."
                        ),
                        "shift_id": shift.id,
                        "shift_name": shift.name,
                        "errors": serializer.errors,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            serializers_to_save.append(serializer)

        created_rules = []

        for serializer in serializers_to_save:
            rule = serializer.save(
                created_by=request.user,
                updated_by=request.user,
            )
            created_rules.append(rule)

        result_serializer = WeeklyOffRuleSerializer(
            created_rules,
            many=True,
            context={"organization": organization},
        )

        return Response(
            {
                "status": True,
                "message": (
                    "Weekly-off rule created for all active shifts."
                ),
                "data": result_serializer.data,
            },
            status=status.HTTP_201_CREATED,
        )


# =========================================================
# WEEKLY OFF RULE UPDATE
# =========================================================

class WeeklyOffRuleUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        rule_id = request.data.get("rule_id")

        if not rule_id:
            return Response(
                {
                    "status": False,
                    "message": "rule_id is required.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        rule = WeeklyOffRule.objects.filter(
            id=rule_id,
            shift__organization=organization,
            is_active=True,
        ).select_related("shift").first()

        if not rule:
            return Response(
                {
                    "status": False,
                    "message": "Weekly-off rule not found.",
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        old_data = {
            "shift_id": rule.shift_id,
            "day_of_week": rule.day_of_week,
            "week_of_month": rule.week_of_month,
            "effective_from": str(rule.effective_from),
            "effective_to": (
                str(rule.effective_to)
                if rule.effective_to else None
            ),
        }

        data = request.data.copy()
        data.pop("rule_id", None)
        data.pop("organization_id", None)

        serializer = WeeklyOffRuleSerializer(
            rule,
            data=data,
            partial=True,
            context={"organization": organization},
        )

        if not serializer.is_valid():
            return Response(
                {
                    "status": False,
                    "message": "Validation failed.",
                    "errors": serializer.errors,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        rule = serializer.save(updated_by=request.user)

        new_data = {
            "shift_id": rule.shift_id,
            "day_of_week": rule.day_of_week,
            "week_of_month": rule.week_of_month,
            "effective_from": str(rule.effective_from),
            "effective_to": (
                str(rule.effective_to)
                if rule.effective_to else None
            ),
        }

        log_activity(
            request=request,
            action="UPDATE",
            module="CALENDAR",
            object_type="WeeklyOffRule",
            object_id=rule.id,
            description=(
                f"Weekly-off rule updated for shift {rule.shift.name}."
            ),
            old_data=old_data,
            new_data=new_data,
        )

        return Response(
            {
                "status": True,
                "message": "Weekly-off rule updated successfully.",
                "data": WeeklyOffRuleSerializer(
                    rule,
                    context={"organization": organization},
                ).data,
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# WEEKLY OFF RULE DEACTIVATE
# =========================================================

class WeeklyOffRuleDeactivateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        rule_id = request.data.get("rule_id")

        if not rule_id:
            return Response(
                {
                    "status": False,
                    "message": "rule_id is required.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        rule = WeeklyOffRule.objects.filter(
            id=rule_id,
            shift__organization=organization,
            is_active=True,
        ).select_related("shift").first()

        if not rule:
            return Response(
                {
                    "status": False,
                    "message": "Weekly-off rule not found.",
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        # Close the rule instead of disabling it, so historical dates
        # can still resolve against this rule.
        effective_to_value = request.data.get("effective_to")

        if not effective_to_value:
            return Response(
                {
                    "status": False,
                    "message": (
                        "effective_to is required to close a weekly-off "
                        "rule while preserving historical calendar data."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        effective_to = parse_date(str(effective_to_value))

        if not effective_to:
            return Response(
                {
                    "status": False,
                    "message": "effective_to must use YYYY-MM-DD format.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if effective_to < rule.effective_from:
            return Response(
                {
                    "status": False,
                    "message": (
                        "effective_to cannot be before effective_from."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        old_data = {
            "effective_to": (
                str(rule.effective_to)
                if rule.effective_to else None
            ),
        }

        rule.effective_to = effective_to
        rule.updated_by = request.user
        rule.save(
            update_fields=[
                "effective_to",
                "updated_by",
                "updated_at",
            ]
        )

        log_activity(
            request=request,
            action="UPDATE",
            module="CALENDAR",
            object_type="WeeklyOffRule",
            object_id=rule.id,
            description=(
                f"Weekly-off rule closed for shift {rule.shift.name}."
            ),
            old_data=old_data,
            new_data={"effective_to": str(effective_to)},
        )

        return Response(
            {
                "status": True,
                "message": "Weekly-off rule closed successfully.",
                "data": {
                    "rule_id": rule.id,
                    "effective_to": str(rule.effective_to),
                },
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# HOLIDAY LIST
# =========================================================

class CalendarHolidayListAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        holidays = CalendarHoliday.objects.filter(
            organization=organization,
            is_active=True,
        )

        date_from = request.data.get("date_from")
        date_to = request.data.get("date_to")
        search = request.data.get("search")

        if date_from:
            date_from = parse_date(str(date_from))

            if not date_from:
                return Response(
                    {
                        "status": False,
                        "message": (
                            "date_from must use YYYY-MM-DD format."
                        ),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            holidays = holidays.filter(date__gte=date_from)

        if date_to:
            date_to = parse_date(str(date_to))

            if not date_to:
                return Response(
                    {
                        "status": False,
                        "message": "date_to must use YYYY-MM-DD format.",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            holidays = holidays.filter(date__lte=date_to)

        if date_from and date_to and date_to < date_from:
            return Response(
                {
                    "status": False,
                    "message:": "date_to cannot be earlier than date_from.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if search:
            search = search.strip()
            holidays = holidays.filter(
                Q(name__icontains=search)
                | Q(description__icontains=search)
            )

        serializer = CalendarHolidaySerializer(
            holidays.order_by("date"),
            many=True,
        )

        return Response(
            {
                "status": True,
                "message": "Holidays fetched successfully.",
                "data": serializer.data,
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# HOLIDAY CREATE
# =========================================================

class CalendarHolidayCreateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        data = request.data.copy()
        data.pop("organization_id", None)

        serializer = CalendarHolidayWriteSerializer(
            data=data,
            context={"organization": organization},
        )

        if not serializer.is_valid():
            return Response(
                {
                    "status": False,
                    "message": "Validation failed.",
                    "errors": serializer.errors,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        holiday = serializer.save(
            organization=organization,
            created_by=request.user,
            updated_by=request.user,
        )

        log_activity(
            request=request,
            action="CREATE",
            module="CALENDAR",
            object_type="CalendarHoliday",
            object_id=holiday.id,
            description=f"Holiday {holiday.name} created.",
            new_data={
                "date": str(holiday.date),
                "name": holiday.name,
                "holiday_type": holiday.holiday_type,
                "pay_type": holiday.pay_type,
                "is_optional": holiday.is_optional,
            },
        )

        return Response(
            {
                "status": True,
                "message": "Holiday created successfully.",
                "data": CalendarHolidaySerializer(holiday).data,
            },
            status=status.HTTP_201_CREATED,
        )


# =========================================================
# HOLIDAY UPDATE
# =========================================================

class CalendarHolidayUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        holiday_id = request.data.get("holiday_id")

        if not holiday_id:
            return Response(
                {
                    "status": False,
                    "message": "holiday_id is required.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        holiday = CalendarHoliday.objects.filter(
            id=holiday_id,
            organization=organization,
            is_active=True,
        ).first()

        if not holiday:
            return Response(
                {
                    "status": False,
                    "message": "Holiday not found.",
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        old_data = {
            "date": str(holiday.date),
            "name": holiday.name,
            "holiday_type": holiday.holiday_type,
            "pay_type": holiday.pay_type,
            "is_optional": holiday.is_optional,
            "description": holiday.description,
        }

        data = request.data.copy()
        data.pop("holiday_id", None)
        data.pop("organization_id", None)

        serializer = CalendarHolidayWriteSerializer(
            holiday,
            data=data,
            partial=True,
            context={"organization": organization},
        )

        if not serializer.is_valid():
            return Response(
                {
                    "status": False,
                    "message": "Validation failed.",
                    "errors": serializer.errors,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        holiday = serializer.save(updated_by=request.user)

        new_data = {
            "date": str(holiday.date),
            "name": holiday.name,
            "holiday_type": holiday.holiday_type,
            "pay_type": holiday.pay_type,
            "is_optional": holiday.is_optional,
            "description": holiday.description,
        }

        log_activity(
            request=request,
            action="UPDATE",
            module="CALENDAR",
            object_type="CalendarHoliday",
            object_id=holiday.id,
            description=f"Holiday {holiday.name} updated.",
            old_data=old_data,
            new_data=new_data,
        )

        return Response(
            {
                "status": True,
                "message": "Holiday updated successfully.",
                "data": CalendarHolidaySerializer(holiday).data,
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# HOLIDAY DEACTIVATE
# =========================================================

class CalendarHolidayDeactivateAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        holiday_id = request.data.get("holiday_id")

        if not holiday_id:
            return Response(
                {
                    "status": False,
                    "message": "holiday_id is required.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        holiday = CalendarHoliday.objects.filter(
            id=holiday_id,
            organization=organization,
            is_active=True,
        ).first()

        if not holiday:
            return Response(
                {
                    "status": False,
                    "message": "Holiday not found.",
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        old_data = {"is_active": holiday.is_active}

        holiday.is_active = False
        holiday.updated_by = request.user
        holiday.save(
            update_fields=["is_active", "updated_by", "updated_at"]
        )

        log_activity(
            request=request,
            action="DELETE",
            module="CALENDAR",
            object_type="CalendarHoliday",
            object_id=holiday.id,
            description=f"Holiday {holiday.name} deactivated.",
            old_data=old_data,
            new_data={"is_active": False},
        )

        return Response(
            {
                "status": True,
                "message": "Holiday deactivated successfully.",
            },
            status=status.HTTP_200_OK,
        )


# =========================================================
# HOLIDAY BULK IMPORT FROM EXCEL
# =========================================================

class CalendarHolidayBulkImportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
    MAX_ROWS = 1000

    REQUIRED_COLUMNS = {
        "date",
        "name",
        "holiday_type",
        "pay_type",
    }

    ALLOWED_COLUMNS = {
        "date",
        "name",
        "holiday_type",
        "pay_type",
        "is_optional",
        "description",
    }

    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        uploaded_file = request.FILES.get("file")

        if not uploaded_file:
            return Response(
                {
                    "status": False,
                    "message": "Please upload an Excel file using the 'file' field.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if uploaded_file.size > self.MAX_FILE_SIZE:
            return Response(
                {
                    "status": False,
                    "message": "File size cannot exceed 5 MB.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not uploaded_file.name.lower().endswith(".xlsx"):
            return Response(
                {
                    "status": False,
                    "message": "Only .xlsx Excel files are supported.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            workbook = load_workbook(
                uploaded_file,
                read_only=True,
                data_only=True,
            )
        except Exception:
            return Response(
                {
                    "status": False,
                    "message": "The uploaded Excel file could not be read.",
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            if "Holidays" in workbook.sheetnames:
                worksheet = workbook["Holidays"]
            else:
                worksheet = workbook.active

            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)

            if not header_row:
                return Response(
                    {
                        "status": False,
                        "message": "The Excel sheet is empty.",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            headers = [
                str(value).strip().lower() if value is not None else ""
                for value in header_row
            ]

            if len(headers) != len(set(headers)):
                return Response(
                    {
                        "status": False,
                        "message": "The Excel file contains duplicate column headers.",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            missing_columns = self.REQUIRED_COLUMNS - set(headers)

            if missing_columns:
                return Response(
                    {
                        "status": False,
                        "message": "Required Excel columns are missing.",
                        "errors": {
                            "missing_columns": sorted(missing_columns),
                        },
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            unknown_columns = set(headers) - self.ALLOWED_COLUMNS

            if unknown_columns:
                return Response(
                    {
                        "status": False,
                        "message": "The Excel file contains unsupported columns.",
                        "errors": {
                            "unknown_columns": sorted(unknown_columns),
                        },
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            column_indexes = {
                name: index
                for index, name in enumerate(headers)
            }

            def get_value(row, column_name):
                index = column_indexes.get(column_name)

                if index is None or index >= len(row):
                    return None

                return row[index]

            def parse_holiday_date(value):
                if isinstance(value, datetime):
                    return value.date()

                if isinstance(value, date):
                    return value

                if isinstance(value, str):
                    return parse_date(value.strip())

                return None

            def parse_optional(value):
                if value is None or value == "":
                    return False

                if isinstance(value, bool):
                    return value

                if isinstance(value, int) and value in (0, 1):
                    return bool(value)

                if isinstance(value, str):
                    normalized = value.strip().lower()

                    if normalized in {"true", "yes", "1"}:
                        return True

                    if normalized in {"false", "no", "0"}:
                        return False

                raise ValueError(
                    "is_optional must be TRUE or FALSE."
                )

            prepared_rows = []
            errors = []
            seen_dates = set()

            for row_number, row in enumerate(rows, start=2):
                if not row or all(
                    value is None or (
                        isinstance(value, str)
                        and not value.strip()
                    )
                    for value in row
                ):
                    continue

                if len(prepared_rows) >= self.MAX_ROWS:
                    errors.append({
                        "row": row_number,
                        "errors": (
                            f"Maximum {self.MAX_ROWS} holiday rows "
                            "are allowed per import."
                        ),
                    })
                    break

                holiday_date = parse_holiday_date(
                    get_value(row, "date")
                )

                if holiday_date is None:
                    errors.append({
                        "row": row_number,
                        "errors": {
                            "date": "Enter a valid date in YYYY-MM-DD format."
                        },
                    })
                    continue

                if holiday_date in seen_dates:
                    errors.append({
                        "row": row_number,
                        "errors": {
                            "date": (
                                "This date appears more than once "
                                "in the uploaded file."
                            )
                        },
                    })
                    continue

                seen_dates.add(holiday_date)

                raw_name = get_value(row, "name")
                name = (
                    str(raw_name).strip()
                    if raw_name is not None
                    else ""
                )

                raw_description = get_value(row, "description")
                description = (
                    str(raw_description).strip()
                    if raw_description is not None
                    else ""
                )

                try:
                    is_optional = parse_optional(
                        get_value(row, "is_optional")
                    )
                except ValueError as exc:
                    errors.append({
                        "row": row_number,
                        "errors": {
                            "is_optional": str(exc),
                        },
                    })
                    continue

                data = {
                    "date": holiday_date,
                    "name": name,
                    "holiday_type": (
                        str(get_value(row, "holiday_type") or "")
                        .strip()
                        .upper()
                    ),
                    "pay_type": (
                        str(get_value(row, "pay_type") or "")
                        .strip()
                        .upper()
                    ),
                    "is_optional": is_optional,
                    "description": description,
                }

                serializer = CalendarHolidayWriteSerializer(
                    data=data,
                    context={
                        "organization": organization,
                    },
                )

                if not serializer.is_valid():
                    errors.append({
                        "row": row_number,
                        "errors": serializer.errors,
                    })
                    continue

                prepared_rows.append({
                    "row_number": row_number,
                    "serializer": serializer,
                })

            if not prepared_rows and not errors:
                errors.append({
                    "row": None,
                    "errors": "The Excel file contains no holiday records.",
                })

            # Validate the entire file before writing anything.
            if errors:
                return Response(
                    {
                        "status": False,
                        "message": (
                            "Holiday import failed. No holidays were imported."
                        ),
                        "errors": errors,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            created_holidays = []

            try:
                # The nested transaction lets us catch database conflicts
                # and return a response without partially saving the file.
                with transaction.atomic():
                    for item in prepared_rows:
                        holiday = item["serializer"].save(
                            organization=organization,
                            created_by=request.user,
                            updated_by=request.user,
                        )

                        created_holidays.append(holiday)

                        log_activity(
                            request=request,
                            action="CREATE",
                            module="CALENDAR",
                            object_type="CalendarHoliday",
                            object_id=holiday.id,
                            description=(
                                f"Holiday {holiday.name} created "
                                "through Excel import."
                            ),
                            new_data={
                                "date": str(holiday.date),
                                "name": holiday.name,
                                "holiday_type": holiday.holiday_type,
                                "pay_type": holiday.pay_type,
                                "is_optional": holiday.is_optional,
                            },
                        )

            except IntegrityError:
                return Response(
                    {
                        "status": False,
                        "message": (
                            "Import failed because one or more holiday "
                            "dates already exist. No holidays were imported."
                        ),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            return Response(
                {
                    "status": True,
                    "message": "Holidays imported successfully.",
                    "data": {
                        "created_count": len(created_holidays),
                        "holidays": CalendarHolidaySerializer(
                            created_holidays,
                            many=True,
                        ).data,
                    },
                },
                status=status.HTTP_201_CREATED,
            )

        finally:
            workbook.close()


# =========================================================
# EMPLOYEE COMBINED CALENDAR
# =========================================================

class EmployeeCalendarAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        organization, error_response = get_user_organization(request)

        if error_response:
            return error_response

        employee = get_employee_for_user(
            request.user,
            organization,
        )

        if not employee:
            return Response({
                "status": False,
                "message": "No active employee record found for this organization.",
                "data": None,
            }, status=status.HTTP_404_NOT_FOUND)

        request_serializer = EmployeeCalendarRequestSerializer(
            data=request.data
        )

        if not request_serializer.is_valid():
            return Response({
                "status": False,
                "message": "Invalid request data.",
                "errors": request_serializer.errors,
            }, status=status.HTTP_400_BAD_REQUEST)

        date_from = request_serializer.validated_data["date_from"]
        date_to = request_serializer.validated_data["date_to"]

        calendar_data = get_employee_calendar(
            organization=organization,
            employee=employee,
            date_from=date_from,
            date_to=date_to,
        )

        calendar_serializer = EmployeeCalendarDaySerializer(
            calendar_data,
            many=True,
        )

        return Response({
            "status": True,
            "message": "Employee calendar fetched successfully.",
            "data": {
                "employee_id": employee.id,
                "date_from": date_from,
                "date_to": date_to,
                "calendar": calendar_serializer.data,
            },
        }, status=status.HTTP_200_OK)
